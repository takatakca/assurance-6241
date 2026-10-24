"""Build a local searchable corpus without changing a source document.

Read-only inputs: the three dossier trees and master/04_Remise_et_suivis.
Outputs: corpus.json and the extraction-coverage JSON/Markdown. OCR renders/cache
live under scripts/.build-livrables/master-corpus-cache, excluded from publication.
No network call, external OCR API, e-mail mutation or publication is performed.
"""
from __future__ import annotations
import argparse, collections, copy, datetime, email, hashlib, html, importlib.util
import json, logging, os, re, subprocess, sys, tempfile, time, zipfile, unicodedata
from pathlib import Path
from html.parser import HTMLParser
from email import policy
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
MASTER = ROOT / '00_DOSSIER_MAITRE'
CORPUS = MASTER / '06_Index_recherche/corpus.json'
COVER = MASTER / '07_Controle_des_sources/COUVERTURE-EXTRACTION.json'
CASES = ['6241 Léger', '6280 Somerled NDG', 'HARMONY MOVERS']
FOLLOWUP = '00_DOSSIER_MAITRE/04_Remise_et_suivis'
RUNTIME = Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies'
POPPLER = RUNTIME/'native/poppler/Library/bin/pdftoppm.exe'
CACHE = ROOT/'scripts/.build-livrables/master-corpus-cache'
SUPPORTED = {'.pdf','.png','.jpg','.jpeg','.webp','.gif','.docx','.xlsx','.xls','.txt','.md','.eml','.csv','.html','.doc','.pptx'}
IMAGES = {'.png','.jpg','.jpeg','.webp','.gif'}
DERIVED_DIRS = {'Rendus','Rotations','Texte_et_OCR','Texte_et_OCR_reprise','Textes_PDF'}
TECH_NAMES = {'sources-mime.json','recherches-gmail.json','couverture.json','provenance.json'}
VERSION = 1
PREVIOUS_DOCS = {}
logging.getLogger('pypdf').setLevel(logging.ERROR)

def read_json(p, default=None):
    try:return json.loads(p.read_text(encoding='utf-8-sig'))
    except Exception:return default

def sha(p):
    h=hashlib.sha256()
    with p.open('rb') as f:
        for b in iter(lambda:f.read(1024*1024),b''):h.update(b)
    return h.hexdigest()

def rel(p):return p.relative_to(ROOT).as_posix()

def atomic_text(p, value):
    p.parent.mkdir(parents=True,exist_ok=True)
    with tempfile.NamedTemporaryFile(mode='w',encoding='utf-8',dir=p.parent,prefix='.'+p.name+'.',suffix='.tmp',delete=False) as f:
        f.write(value);temporary=Path(f.name)
    temporary.replace(p)

def local_path(value, base=None):
    s=str(value).replace('\\','/')
    s=re.sub('/+','/',s)
    p=Path(s)
    if not p.is_absolute():p=(base or ROOT)/p
    try:
        p=p.resolve()
        if p.is_relative_to(ROOT.resolve()) and p.is_file():return p
    except OSError:pass
    return None

def text_file(p):
    b=p.read_bytes()
    for enc in ('utf-8-sig','utf-16','cp1252'):
        try:return b.decode(enc)
        except (UnicodeDecodeError,UnicodeError):pass
    return b.decode('utf-8',errors='replace')

class HTMLText(HTMLParser):
    def __init__(self):super().__init__(convert_charrefs=True);self.bits=[];self.hidden=0
    def handle_starttag(self,tag,attrs):
        if tag in {'script','style','noscript'}:self.hidden+=1
        if tag in {'p','br','div','tr','li','h1','h2','h3','h4'}:self.bits.append('\n')
    def handle_endtag(self,tag):
        if tag in {'script','style','noscript'}:self.hidden=max(0,self.hidden-1)
        if tag in {'p','div','tr','li'}:self.bits.append('\n')
    def handle_data(self,data):
        if not self.hidden:self.bits.append(data)

def html_text(t):
    h=HTMLText();h.feed(t);return ''.join(h.bits)

def load_remote_rules():
    spec=importlib.util.spec_from_file_location('_master_remote_rules',ROOT/'scripts/build_remote_site.py')
    mod=importlib.util.module_from_spec(spec);spec.loader.exec_module(mod)
    return mod

def redact(t, secret_patterns):
    counts=collections.Counter()
    for label,pattern in secret_patterns.items():
        rx=pattern.decode('ascii')
        t,n=re.subn(rx,'[IDENTIFIANT DE SERVICE MASQUÉ]',t);counts[label]+=n
    # A credential value can be on the next line of a MIME or OCR extraction.
    credential=re.compile(r'(?i)(?:mot\s+de\s+passe|password|passcode|username|user\s+name|nom\s+d.utilisateur|login\s+(?:id|information)|code\s+(?:secret|d.accès)|security\s+(?:answer|question)|réponse\s+de\s+sécurité)')
    lines=t.splitlines();out=[];skip=0
    for line in lines:
        if skip and line.strip():
            out.append('[CONTEXTE D’ACCÈS MASQUÉ]');skip-=1;counts['credential_context']+=1;continue
        if credential.search(line) and 'MASQU' not in line.upper():
            out.append('[MENTION OU VALEUR D’ACCÈS MASQUÉE]');skip=2;counts['credential_line']+=1
        else:out.append(line)
    t='\n'.join(out)
    # Avoid leaking onboarding links and URL-embedded access tokens into search.
    def url(m):
        u=m.group(0)
        if any(v in u.lower() for v in ('connect.stripe.com/setup/','mail.google.com/mail/u/','urldefense','links.notification.intuit','url9448.','click.pstmrk')) or re.search(r'(?i)[?&](token|access_token|auth|password|key|secret|signature|sig|code)=',u):
            counts['access_url']+=1;return '[LIEN TECHNIQUE D’ACCÈS MASQUÉ]'
        return u
    t=re.sub(r'https?://[^\s<>"\)]+',url,t)
    return t,dict((k,v) for k,v in counts.items() if v)

def source_type(p):
    s=rel(p).lower()
    ext=p.suffix.lower()
    if s.startswith(FOLLOWUP.lower()+'/'):return 'projet_de_remise_non_envoye'
    if ext=='.md' or any(w in p.stem.lower() for w in ('reponses-aux-25-demandes','synthese','synthèse','bilan')):return 'analyse_ou_document_de_travail'
    # A scan in an analysis folder is still a source document, not the analysis.
    if ext in {'.txt','.html','.csv'} and any(w in s for w in ('lectures/','texte','transcription','ocr','extraction')):return 'extraction_ou_transcription'
    return 'document_source_ou_copie'

SOURCE_LABELS={
    'document_source_ou_copie':'Document source ou copie — attribution à vérifier sur la pièce',
    'extraction_ou_transcription':'Extraction ou transcription — à rapprocher de la pièce',
    'analyse_ou_document_de_travail':'Analyse ou document de travail — ne vaut pas pièce originale',
    'projet_de_remise_non_envoye':'Projet de remise ou suivi — document de travail, envoi non établi',
}

def select_files(rules):
    selected=[];excluded=[];all_files=[]
    roots=[ROOT/n for n in CASES]
    if (ROOT/FOLLOWUP).is_dir():roots.append(ROOT/FOLLOWUP)
    for base in roots:
        for p in sorted(base.rglob('*')):
            if not p.is_file():continue
            all_files.append(p);s=rel(p);parts=p.relative_to(ROOT).parts;reason=None
            if p.is_symlink() or not p.resolve().is_relative_to(ROOT.resolve()):reason='lien_hors_perimetre'
            elif any(x in rules.EXCLUDE_PARTS for x in parts) or s.startswith(rules.LOCAL_COLLECTION_PREFIXES):reason='exclusion_identique_site_collecte_technique'
            elif any(x.upper() in {'RAW','MIME','RAW_TRANSPORT'} for x in parts):reason='export_technique_brut_non_indexe'
            elif p.name.lower() in rules.SECRET_NAMES or p.suffix.lower() in {'.pem','.key','.pfx','.p12','.b64url'}:reason='secret_ou_transport_technique'
            elif s.startswith(FOLLOWUP+'/') and p.suffix.lower() not in {'.md','.txt'}:reason='master_hors_MD_TXT_autorises'
            elif any(x in DERIVED_DIRS for x in parts):reason='derive_texte_ocr_rendu_rattache_aux_originaux'
            elif p.suffix.lower()=='.json':reason='index_ou_export_JSON_technique_non_corpus'
            elif p.suffix.lower()=='.zip':reason='archive_conteneur_non_decompressee_contenu_interne_non_audite'
            elif p.suffix.lower() not in SUPPORTED:reason='format_non_pris_en_charge'
            if reason:excluded.append({'path':s,'reason':reason,'bytes':p.stat().st_size})
            else:selected.append(p)
    return selected,excluded,all_files

def preference(p):
    s=rel(p);return ('ANNEXES' in p.parts,'12_Remise_Dufourd' in s,source_type(p)=='extraction_ou_transcription',len(p.parts),len(s),s)

def normalized_body(t):
    t=unicodedata.normalize('NFKD',html.unescape(t)).casefold()
    t=''.join(c for c in t if not unicodedata.combining(c))
    return ' '.join(re.findall(r'\w+',t))

def header_text(value):
    if value is None:return ''
    if isinstance(value,str):return value
    if isinstance(value,list):return '; '.join(header_text(x) for x in value)
    if isinstance(value,dict):return ' '.join(str(value[k]) for k in ('name','email','address') if value.get(k))
    return str(value)

def email_json_messages(p):
    """Read only known mail exports, never search results, ChatGPT or binaries."""
    s=rel(p);parts=p.relative_to(ROOT).parts
    if len(parts)<3 or parts[0] not in CASES[:2] or parts[1]!='03_Courriels':return []
    if any(x in DERIVED_DIRS or x in {'99_Travail_et_collecte','recherches'} for x in parts):return []
    if not (p.name=='emails-index.json' or re.fullmatch(r'[0-9a-f]{16}\.json',p.name) or re.fullmatch(r'fils_(?:bruts|complementaires).*\.json',p.name)):return []
    obj=read_json(p)
    rows=[]
    def visit(x):
        if isinstance(x,list):
            for y in x:visit(y)
        elif isinstance(x,dict):
            if x.get('id') and (isinstance(x.get('body'),str) or isinstance(x.get('payload'),dict)):
                rows.append(x);return
            if p.name=='emails-index.json' and x is obj:
                visit(x.get('messages',[]));visit(x.get('additional_messages',[]));return
            for k,v in x.items():
                if k not in {'content','body','base64_url_content','data','attachments','headers'} and isinstance(v,(dict,list)):visit(v)
    visit(obj)
    output=[]
    for row in rows:
        attachments=[];limitations=[]
        def mime(part):
            if part.get('filename'):
                attachments.append({'filename':str(part['filename']),'mime':str(part.get('mime_type',part.get('mimeType','')))})
                return ''
            typ=str(part.get('mime_type',part.get('mimeType',''))).lower()
            children=part.get('parts') or []
            if children:
                if typ=='multipart/alternative':
                    plain=[x for x in children if str(x.get('mime_type',x.get('mimeType',''))).lower()=='text/plain']
                    chosen=plain or [x for x in children if str(x.get('mime_type',x.get('mimeType',''))).lower()=='text/html'] or children[:1]
                else:chosen=children
                return '\n'.join(mime(x) for x in chosen)
            body=part.get('body') or {}
            content=body.get('content','') if isinstance(body,dict) else ''
            if typ in {'text/plain','text/html'} and isinstance(content,str):
                if not content and body.get('base64_url_content'):limitations.append('Partie texte encodée non décodée dans cette extraction; aucun base64 indexé.')
                return html_text(content) if typ=='text/html' else content
            return ''
        if isinstance(row.get('body'),str):
            body=row['body']
            headers={k:header_text(row.get(k,'')) for k in ('from','to','cc','date','subject','rfc_message_id')}
            if not headers['date']:headers['date']=header_text(row.get('date_iso',''))
            for a in row.get('attachments',[]):
                if isinstance(a,dict) and a.get('filename'):attachments.append({'filename':str(a['filename']),'mime':str(a.get('mime',a.get('mime_type','')))})
        else:
            payload=row['payload'];hm={str(h.get('name','')).lower():header_text(h.get('value','')) for h in payload.get('headers',[]) if isinstance(h,dict)}
            body=mime(payload);headers={k:hm.get(k.replace('rfc_message_id','message-id'),'') for k in ('from','to','cc','date','subject','rfc_message_id')}
        output.append({'gmail_id':str(row['id']),'thread_id':str(row.get('thread_id',row.get('threadId',''))),'headers':headers,'body':body,'attachments':attachments,'limitations':limitations,'path':s})
    return output

def integrate_email_exports(docs,all_files,excluded,rules):
    occurrences=collections.defaultdict(list);file_hashes={};export_paths=set()
    for p in all_files:
        if p.suffix.lower()!='.json':continue
        messages=email_json_messages(p)
        if not messages:continue
        file_hashes[rel(p)]=sha(p);export_paths.add(rel(p))
        for m in messages:occurrences[(rel(p).split('/')[0],m['gmail_id'])].append(m)
    normal_docs=[]
    for d in docs:
        searchtext='\n'.join(pg.get('text','') for pg in d['pages'])
        if searchtext.strip():normal_docs.append((d,normalized_body(redact(searchtext,rules.SECRET_PATTERNS)[0]),(' '.join([d['path']]+d['aliases'])+' '+searchtext).casefold()))
    grouped=collections.defaultdict(list);audit=[];covered=0
    for (case,ident),variants in sorted(occurrences.items()):
        # Preserve every source reference; choose the longest readable body only
        # for search. This does not replace or certify the original export.
        m=max(variants,key=lambda x:(len(x['body']),x['path'].endswith('emails-index.json')))
        clean_body,body_redactions=redact(m['body'],rules.SECRET_PATTERNS)
        needle=normalized_body(clean_body);rfc=m['headers'].get('rfc_message_id','').casefold()
        complete=[]
        if needle:
            for d,normal,locator in normal_docs:
                if (ident.casefold() in locator or (rfc and rfc in locator)) and needle in normal:complete.append(d)
        sources=[{'path':s,'sha256':file_hashes[s]} for s in sorted({x['path'] for x in variants})]
        record={'case':case,'gmail_message_id':ident,'source_exports':sources,'body_characters':len(m['body']),'attachments_announced':len(m['attachments'])}
        if complete:
            target=min(complete,key=lambda d:(d['format']!='EML',len(d['path'])))
            record.update(status='corps_complet_deja_indexe_meme_identifiant',covered_by={'id':target['id'],'path':target['path'],'sha256':target['sha256']})
            target.setdefault('email_export_references',[]).append(record)
            covered+=1
        else:
            headers=m['headers'];lines=['Courriel — export Gmail','Identifiant Gmail : '+ident]
            for label,key in [('Date','date'),('De','from'),('À','to'),('Cc','cc'),('Objet','subject'),('Message-ID','rfc_message_id')]:
                if headers.get(key):lines.append(label+' : '+headers[key])
            lines+=['','CORPS DU COURRIEL',m['body'] or '[Corps lisible absent de cet export]']
            if m['attachments']:
                lines+=['','PIÈCES JOINTES ANNONCÉES — noms seulement; contenu binaire non indexé depuis cet export']
                lines += ['- '+a['filename']+(' ('+a['mime']+')' if a['mime'] else '') for a in m['attachments']]
            lines+=['','PROVENANCE : '+m['path'],'Les citations de messages antérieurs restent dans le corps; elles ne constituent pas de nouveaux envois.']+m['limitations']
            pg={'page':0,'label':headers.get('subject','Sans objet')+' — '+headers.get('date','Date non renseignée')+' — '+ident,'text':'\n'.join(lines),'method':'Gmail_JSON_entetes_corps_lisible_noms_PJ_sans_binaire','gmail_message_id':ident,'source_exports':sources,'extraction_limitations':m['limitations']}
            grouped[m['path']].append(pg)
            record.update(status='corps_ajoute_depuis_export_JSON',indexed_path=m['path'])
        audit.append(record)
    added=[]
    for path,pages in sorted(grouped.items()):
        p=ROOT/path;h=file_hashes[path];case=path.split('/')[0]
        for n,pg in enumerate(pages,1):pg['page']=n
        title=('Courriels '+case+' — '+str(len(pages))+' échanges exportés') if len(pages)>1 else 'Courriel — '+pages[0]['label']
        d={'id':'DOC-'+h[:16].upper(),'sha256':h,'case':case,'cases':[case],'title':title,'path':path,'aliases':[],'format':'JSON','bytes':p.stat().st_size,'source_nature_code':'extraction_ou_transcription','source_nature':'Extraction lisible de courriels Gmail — export source conservé, pièces jointes binaires non décodées','pages':pages,'text':'','extraction_status':'non_traite','errors':[],'attachments_announced':[],'email_export_message_count':len(pages)}
        docs.append(d);added.append(path)
    for row in excluded:
        if row['path'] in export_paths:row['reason']='courriel_JSON_traite_deduplication_par_identifiant'
    excluded[:]=[e for e in excluded if e['path'] not in added]
    return {'export_files_recognized':len(export_paths),'unique_messages':len(occurrences),'messages_added':sum(len(v) for v in grouped.values()),'messages_already_complete':covered,'source_documents_added':len(added),'by_case':dict(collections.Counter(k[0] for k in occurrences)),'messages':audit,'matching_rule':'Même identifiant Gmail ou Message-ID RFC et corps complet normalisé présent; aucune assimilation sur de simples extraits. Les références JSON restent séparées des alias SHA.'},len(added)

def existing_ocr(all_files, groups):
    by_prefix={h[:12].upper():h for h in groups}
    stems=collections.defaultdict(set)
    for h,ps in groups.items():
        for p in ps:
            if p.suffix.lower() in IMAGES|{'.pdf'}:stems[p.stem].add(h)
    results=collections.defaultdict(list);used=set();warnings=[]
    def add(h,page,p,why):
        if not p or not h or not p.exists():return
        t=text_file(p).strip()
        if t:results[(h,int(page))].append({'text':t,'source_ocr':rel(p),'selection':why});used.add(rel(p))
    for p in all_files:
        if p.name=='manifest-extraction.json':
            rows=read_json(p,[])
            if not isinstance(rows,list):continue
            for d in rows:
                h=d.get('sha256')
                if h not in groups:continue
                for page in d.get('pages',[]):
                    n=page.get('page',1)
                    fn=page.get('ocr_file') or f'{d.get("id","DOC-"+h[:12].upper())}-p{n:03}-ocr.txt'
                    add(h,n,local_path(fn,p.parent),'manifest_sha256')
        if p.name in {'ocr-jobs.json','ocr-orientation-selection.json'}:
            rows=read_json(p,[])
            if not isinstance(rows,list):continue
            for d in rows:
                if d.get('error'):continue
                ident=re.sub(r'-rot(?:90|180|270)$','',str(d.get('id','')))
                candidates=set()
                if ident.startswith('DOC-') and ident[4:].upper() in by_prefix:candidates={by_prefix[ident[4:].upper()]}
                else:
                    ident=re.sub(r'-p\d+$','',ident)
                    candidates=stems.get(ident,set())
                if len(candidates)!=1:continue
                out=local_path(d.get('output',''),ROOT/rel(p).split('/')[0])
                if not out:out=local_path(Path(str(d.get('output','')).replace('\\','/')).name,p.parent)
                add(next(iter(candidates)),d.get('page',1),out,'mapping_nom_unique_et_page; OCR existant non validé intégralement')
    return results,used,warnings

def native_extract(p,h):
    prior=PREVIOUS_DOCS.get(h)
    # The sandbox can discard its temporary directory after each process.
    # Reuse the already masked extraction only after the source SHA has matched.
    # A truncated/error-only record must still be retried from the source.
    if prior and prior.get('pages') and not prior.get('truncated'):
        return {'pages':copy.deepcopy(prior['pages']),'errors':copy.deepcopy(prior.get('errors',[])),'attachments':copy.deepcopy(prior.get('attachments_announced',[]))}
    cache=CACHE/'native'/f'{h}.json';cached=read_json(cache)
    if cached and cached.get('version')==VERSION:return cached['result']
    ext=p.suffix.lower();pages=[];errors=[];attachments=[]
    if ext=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(p)
        if reader.is_encrypted and not reader.decrypt(''):raise ValueError('PDF chiffré, mot de passe requis')
        for i,pg in enumerate(reader.pages,1):
            try:t=pg.extract_text() or ''
            except Exception as e:t='';errors.append({'page':i,'error':type(e).__name__+': '+str(e)[:160]})
            pages.append({'page':i,'text':t,'method':'pdf_texte_natif'})
    elif ext in IMAGES:
        from PIL import Image
        with Image.open(p) as im:
            n=getattr(im,'n_frames',1)
            for i in range(n):pages.append({'page':i+1,'text':'','method':'image_sans_texte_natif'})
    elif ext in {'.txt','.md','.csv'}:
        pages=[{'page':1,'text':text_file(p),'method':'texte_UTF8_ou_decodage_source'}]
    elif ext=='.html':pages=[{'page':1,'text':html_text(text_file(p)),'method':'html_texte_visible_sans_scripts'}]
    elif ext=='.eml':
        msg=email.message_from_bytes(p.read_bytes(),policy=policy.default)
        header='\n'.join(k+': '+str(msg.get(k,'')) for k in ['Date','From','To','Cc','Subject','Message-ID'])
        def body(m):
            if m.get_content_type()=='message/rfc822':return '\n[MESSAGE JOINT]\n'+'\n'.join(body(x) for x in m.get_payload())
            if m.is_multipart():
                children=list(m.iter_parts())
                if m.get_content_subtype()=='alternative':
                    plain=[c for c in children if c.get_content_type()=='text/plain']
                    return '\n'.join(body(c) for c in (plain or children[:1]))
                return '\n'.join(body(c) for c in children)
            if m.get_filename():attachments.append({'filename':m.get_filename(),'mime':m.get_content_type()});return '[Pièce jointe annoncée : '+str(m.get_filename())+']'
            if m.get_content_type() in {'text/plain','text/html'}:
                try:t=m.get_content()
                except Exception:t=(m.get_payload(decode=True) or b'').decode(m.get_content_charset() or 'utf-8',errors='replace')
                return html_text(t) if m.get_content_type()=='text/html' else t
            return ''
        pages=[{'page':1,'text':header+'\n\n'+body(msg),'method':'EML_MIME_corps_et_citations_PJ_non_extractees'}]
    elif ext in {'.docx','.pptx'}:
        with zipfile.ZipFile(p) as z:
            if ext=='.docx':names=[n for n in z.namelist() if re.fullmatch(r'word/(document|header\d+|footer\d+|footnotes|endnotes)\.xml',n)]
            else:names=sorted([n for n in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml',n)],key=lambda n:int(re.search(r'(\d+)\.xml',n)[1]))
            names=sorted(names,key=lambda n:(n!='word/document.xml',n)) if ext=='.docx' else names
            for i,n in enumerate(names,1):
                root=ET.fromstring(z.read(n));chunks=[]
                for para in root.iter():
                    if para.tag.rsplit('}',1)[-1]=='p':chunks.append(''.join(x.text or '' for x in para.iter() if x.tag.rsplit('}',1)[-1]=='t'))
                pages.append({'page':i,'label':n,'text':'\n'.join(chunks),'method':'OOXML_paragraphes_tables; unités_logiques_non_pages_imprimees'})
    elif ext=='.xlsx':
        import openpyxl
        wb=openpyxl.load_workbook(p,read_only=True,data_only=False);cached_values=openpyxl.load_workbook(p,read_only=True,data_only=True)
        for i,ws in enumerate(wb.worksheets,1):
            lines=['Feuille : '+ws.title];cv=cached_values[ws.title]
            for rn,row in enumerate(ws.iter_rows(),1):
                cells=[]
                for cell in row:
                    value=cell.value
                    if value is None:continue
                    if cell.data_type=='f':value=str(value)+' [valeur en cache : '+str(cv[cell.coordinate].value)+']'
                    cells.append(str(cell.column_letter)+'='+str(value))
                if cells:lines.append('Ligne '+str(rn)+' : '+' | '.join(cells))
            pages.append({'page':i,'label':ws.title,'text':'\n'.join(lines),'method':'xlsx_cellules_et_formules_cachees_non_recalculees'})
        wb.close();cached_values.close()
    elif ext=='.xls':
        try:import xlrd
        except ImportError:raise ValueError('XLS ancien : lecteur xlrd indisponible; original préservé, non indexé')
        wb=xlrd.open_workbook(p)
        for i,s in enumerate(wb.sheets(),1):
            pages.append({'page':i,'label':s.name,'text':'Feuille : '+s.name+'\n'+'\n'.join('Ligne '+str(r+1)+' : '+' | '.join(str(v) for v in s.row_values(r))for r in range(s.nrows)),'method':'xls_valeurs_enregistrees'})
    elif ext=='.doc':raise ValueError('DOC binaire ancien : convertisseur local non disponible; original préservé')
    result={'pages':pages,'errors':errors,'attachments':attachments}
    cache.parent.mkdir(parents=True,exist_ok=True);cache.write_text(json.dumps({'version':VERSION,'result':result},ensure_ascii=False),encoding='utf-8')
    return result

def run_ocr(pending,limit):
    from PIL import Image,ImageOps
    work=CACHE/'ocr';work.mkdir(parents=True,exist_ok=True)
    jobs=[];mapped={};results={}
    for d,pg,p in pending:
        key=(d['sha256'],pg['page']);stem=d['sha256']+'-p'+str(pg['page']);dest=work/(stem+'.txt')
        if dest.exists():
            t=text_file(dest)
            if t.strip():results[key]={'text':t,'method':'ocr_Windows_local_a_controler','source_ocr':'cache_local_exclu_publication','validation_visuelle':'non_effectuee'}
            continue
        if len(jobs)>=limit:continue
        if p.suffix.lower()=='.pdf' and not POPPLER.exists():continue
        image_path=work/(stem+'.png')
        try:
            if p.suffix.lower()=='.pdf':
                proc=subprocess.run([str(POPPLER),'-f',str(pg['page']),'-l',str(pg['page']),'-singlefile','-scale-to','2400','-png',str(p),str(image_path.with_suffix(''))],capture_output=True,timeout=40)
                if proc.returncode:continue
            else:
                with Image.open(p) as im:
                    im.seek(pg['page']-1);im=ImageOps.exif_transpose(im).convert('RGB');im.thumbnail((2400,2400));im.save(image_path)
            jobs.append({'id':d['sha256'],'page':pg['page'],'image':str(image_path),'output':str(dest)});mapped[key]=dest
        except Exception:continue
    if jobs:
        jobfile=work/'ocr-jobs.json';jobfile.write_text(json.dumps(jobs,ensure_ascii=False),encoding='utf-8-sig')
        ps=Path(os.environ.get('WINDIR','C:/Windows'))/'System32/WindowsPowerShell/v1.0/powershell.exe'
        try:
            proc=subprocess.run([str(ps),'-NoProfile','-ExecutionPolicy','Bypass','-File',str(ROOT/'scripts/run_case_ocr.ps1'),'-JobsPath',str(jobfile)],capture_output=True,timeout=max(90,len(jobs)*6))
            state='Windows OCR local; '+str(len(jobs))+' pages demandées; code '+str(proc.returncode)
        except Exception as e:state='OCR local interrompu/indisponible : '+type(e).__name__
        for key,dest in mapped.items():
            if dest.exists() and text_file(dest).strip():results[key]={'text':text_file(dest),'method':'ocr_Windows_local_a_controler','source_ocr':'cache_local_exclu_publication','validation_visuelle':'non_effectuee'}
    else:state='Aucune nouvelle page OCR demandée; cache utilisé si disponible'
    return results,state

def main():
    global PREVIOUS_DOCS
    ap=argparse.ArgumentParser(description=__doc__);ap.add_argument('--ocr-limit',type=int,default=0);ap.add_argument('--max-text-chars',type=int,default=400000);args=ap.parse_args()
    started=time.monotonic();rules=load_remote_rules();selected,excluded,all_files=select_files(rules)
    PREVIOUS_DOCS={d['sha256']:d for d in (read_json(CORPUS,{}) or {}).get('documents',[])}
    groups=collections.defaultdict(list)
    for i,p in enumerate(selected,1):
        try:groups[sha(p)].append(p)
        except OSError as e:excluded.append({'path':rel(p),'reason':'lecture_fichier_impossible','error':str(e)[:200]})
        if i%250==0:print(f'Empreintes : {i}/{len(selected)}',flush=True)
    ocr_map,ocr_used,warnings=existing_ocr(all_files,groups)
    docs=[];pending=[];duplicates=[];errors=[]
    for i,(h,ps) in enumerate(sorted(groups.items(),key=lambda x:preference(min(x[1],key=preference))),1):
        ps=sorted(ps,key=preference);p=ps[0];aliases=[rel(x) for x in ps[1:]]
        case='Remise et suivis' if rel(p).startswith(FOLLOWUP+'/') else rel(p).split('/')[0]
        d={'id':'DOC-'+h[:16].upper(),'sha256':h,'case':case,'cases':sorted({'Remise et suivis' if rel(x).startswith(FOLLOWUP+'/') else rel(x).split('/')[0] for x in ps}),'title':p.name,'path':rel(p),'aliases':aliases,'format':p.suffix[1:].upper(),'bytes':p.stat().st_size,'source_nature_code':source_type(p),'source_nature':SOURCE_LABELS[source_type(p)],'pages':[],'text':'','extraction_status':'non_traite','errors':[]}
        if aliases:duplicates.append({'sha256':h,'canonical_path':d['path'],'aliases':aliases})
        try:r=native_extract(p,h);d['pages']=r['pages'];d['errors']=r['errors'];d['attachments_announced']=r['attachments']
        except Exception as e:d['errors'].append({'error':type(e).__name__+': '+str(e)[:220]});d['extraction_status']='non_indexe_erreur_ou_format'
        for pg in d['pages']:
            if p.suffix.lower() in IMAGES|{'.pdf'} and len(re.sub(r'\s','',pg['text']))<100:
                candidates=ocr_map.get((h,pg['page']),[])
                if candidates:
                    chosen=max(candidates,key=lambda x:len(x['text']))
                    if len(chosen['text'].strip())>len(pg['text'].strip()):
                        pg['native_text_characters']=len(pg['text']);pg.update(chosen);pg['method']='ocr_existant_a_controler';pg['validation_visuelle']='non_effectuee_par_cette_indexation'
                if len(re.sub(r'\s','',pg['text']))<100:pending.append((d,pg,p))
        docs.append(d)
        if i%50==0:print(f'Extraction : {i}/{len(groups)} documents',flush=True)
    mail_audit,mail_source_files=integrate_email_exports(docs,all_files,excluded,rules)
    print('Courriels JSON : '+str(mail_audit['unique_messages'])+' messages uniques; '+str(mail_audit['messages_added'])+' ajoutés; '+str(mail_audit['messages_already_complete'])+' déjà complets.',flush=True)
    # Prefer small recent source documents over generated reports for the targeted OCR pass.
    pending.sort(key=lambda x:(x[0]['sha256']!='da3e5f8ebdf17514bf62016707ea6207eb8f2fafa9287b6cbceaaa7b84ab4b1b','14_Recherche' not in x[2].as_posix(),'13_Complements' not in x[2].as_posix(),len(x[0]['pages']),x[0]['path'],x[1]['page']))
    new_ocr,ocr_state=run_ocr(pending,max(0,args.ocr_limit))
    for d,pg,p in pending:
        result=new_ocr.get((d['sha256'],pg['page']))
        if result and len(result['text'].strip())>len(pg['text'].strip()):pg.update(result)
    redactions=[];nonindexed=[];truncated=[];methods=collections.Counter()
    for d in docs:
        prior_redactions={} if d['format']=='JSON' else PREVIOUS_DOCS.get(d['sha256'],{}).get('redactions',{})
        remaining=args.max_text_chars;original_chars=0;parts=[];missing_pages=[];low_text_pages=[];ocr_pages=[];red=collections.Counter(prior_redactions)
        for pg in d['pages']:
            t,counts=redact(pg['text'],rules.SECRET_PATTERNS);red.update(counts);original_chars+=len(t)
            original=len(t)
            if original>remaining:
                t=t[:max(0,remaining)];pg['truncated']=True;pg['characters_available']=original;pg['characters_indexed']=len(t)
                if t:t+='\n[FIN DE L’EXTRAIT INDEXÉ — limite explicite; consulter le document original]'
            remaining=max(0,remaining-min(original,remaining));pg['text']=t;methods[pg['method']]+=1
            if t.strip():parts.append('[Page/unité '+str(pg['page'])+(' — '+pg['label'] if pg.get('label') else '')+'; '+pg['method']+']\n'+t)
            if not t.strip() and not pg.get('truncated'):missing_pages.append(pg['page'])
            if d['format'].lower() in {'pdf','png','jpg','jpeg','webp','gif'} and len(re.sub(r'\s','',t))<100 and 'ocr' not in pg['method'].lower() and not pg.get('truncated'):
                low_text_pages.append(pg['page']);pg['scan_or_low_text_to_verify']=True
            if 'ocr' in pg['method'].lower():ocr_pages.append(pg['page'])
        d['text']='\n\n'.join(parts);d['characters_available_after_redaction']=original_chars;d['characters_indexed']=sum(len(p['text']) for p in d['pages']);d['truncated']=any(p.get('truncated') for p in d['pages']);d['missing_text_pages']=missing_pages;d['low_text_pages_to_verify']=low_text_pages;d['ocr_pages_to_verify']=ocr_pages;d['redactions']=dict(red)
        if d['text'].strip():
            d['extraction_status']='partiel_limite_explicite' if d['truncated'] else 'partiel_pages_sans_texte' if missing_pages or d['errors'] else 'partiel_texte_faible_a_controler' if low_text_pages else 'texte_disponible_ocr_a_controler' if ocr_pages else 'texte_natif_disponible'
        elif d['extraction_status']=='non_traite':d['extraction_status']='non_indexe_aucun_texte_detecte'
        if missing_pages or low_text_pages or not d['text'].strip():nonindexed.append({'id':d['id'],'path':d['path'],'status':d['extraction_status'],'pages_without_text':missing_pages,'low_text_pages_to_verify':low_text_pages,'errors':d['errors']})
        if d['truncated']:truncated.append({'id':d['id'],'path':d['path'],'characters_available':original_chars,'characters_indexed':d['characters_indexed'],'limit':args.max_text_chars})
        if red:redactions.append({'id':d['id'],'path':d['path'],'counts':dict(red)})
    prior_checks=(read_json(COVER,{}) or {}).get('visual_sample_checks',[])
    checks=[]
    docs_by_sha={d['sha256']:d for d in docs}
    for check in prior_checks:
        d=docs_by_sha.get(check.get('sha256'))
        if not d:continue
        check=dict(check);check['canonical_path']=d['path'];checks.append(check)
        d.setdefault('visual_sample_checks',[]).append(check)
        for pg in d['pages']:
            if pg['page']==check.get('page'):
                pg['visual_review_note']=check.get('findings','')
                pg['validation_visuelle']='contrôle ciblé consigné; OCR non corrigé ni validé intégralement'
    stats={'source_files_seen':len(all_files),'source_files_selected':len(selected)+mail_source_files,'documents_unique':len(docs),'duplicate_paths':sum(len(d['aliases']) for d in docs),'excluded_files':len(excluded),'documents_with_text':sum(bool(d['text'].strip()) for d in docs),'documents_without_text':sum(not d['text'].strip() for d in docs),'pages_or_units':sum(len(d['pages']) for d in docs),'pages_without_text':sum(len(d['missing_text_pages']) for d in docs),'documents_truncated':len(truncated),'documents_with_redactions':len(redactions),'ocr_existing_candidates':len(ocr_map),'ocr_new_or_cached_pages':len(new_ocr),'visual_checks':len(checks),'email_json_unique_messages':mail_audit['unique_messages'],'email_json_messages_added':mail_audit['messages_added'],'email_json_messages_already_complete':mail_audit['messages_already_complete'],'by_case':dict(collections.Counter(d['case'] for d in docs)),'by_format':dict(collections.Counter(d['format'] for d in docs)),'by_status':dict(collections.Counter(d['extraction_status'] for d in docs)),'by_method':dict(methods),'elapsed_seconds':round(time.monotonic()-started,1)}
    now=datetime.datetime.now(datetime.timezone.utc).isoformat()
    warnings+=['Les textes servent à la recherche, pas à certifier les faits, les signatures, les montants ou l’exhaustivité des courriels.','Les unités DOCX/XLSX/PPTX sont des parties logiques, feuilles ou diapositives; seules les unités PDF correspondent aux pages imprimées.','OCR automatique à contrôler sur les images originales; le contrôle visuel ciblé ne valide pas toutes les pages.','Pièces jointes EML : noms annoncés indexés, contenu binaire non extrait depuis les EML; les copies classées séparément sont indexées quand disponibles.','Les doublons exacts sont groupés par SHA; les versions distinctes restent distinctes et leurs montants ne sont pas additionnés.','Les JSON techniques, caches et rendus sont exclus; exclusions détaillées dans la couverture. Les archives ZIP ne sont pas décompressées par cet indexeur et leur contenu interne n’est pas audité; les copies classées séparément sont indexées quand disponibles.']
    payload={'generated_at':now,'schema_version':1,'documents':docs,'stats':stats,'warnings':warnings}
    preserved_ocr=[]
    for d in docs:
        old=PREVIOUS_DOCS.get(d['sha256'],{})
        old_ocr={p['page'] for p in old.get('pages',[]) if p.get('text','').strip() and 'ocr' in p.get('method','').lower()}
        new_ocr_pages={p['page'] for p in d['pages'] if p.get('text','').strip() and 'ocr' in p.get('method','').lower()}
        lost=old_ocr-new_ocr_pages
        if lost:raise RuntimeError('Régression OCR refusée pour '+d['sha256']+' pages '+str(sorted(lost)))
        if old_ocr:preserved_ocr.append({'sha256':d['sha256'],'pages':sorted(old_ocr)})
        for pg in d['pages']:
            if pg['method']=='ocr_Windows_local_a_controler' and pg['text'].strip():
                atomic_text(CACHE/'ocr'/(d['sha256']+'-p'+str(pg['page'])+'.txt'),pg['text'])
    stats['ocr_windows_pages_total']=sum(p['method']=='ocr_Windows_local_a_controler' for d in docs for p in d['pages'])
    atomic_text(CORPUS,json.dumps(payload,ensure_ascii=False,separators=(',',':')))
    stats['corpus_bytes']=CORPUS.stat().st_size
    coverage={'generated_at':now,'corpus_path':rel(CORPUS),'corpus_sha256':sha(CORPUS),'stats':stats,'limits':{'max_text_chars_per_document':args.max_text_chars,'new_ocr_job_limit':args.ocr_limit,'ocr_runtime':ocr_state,'network_calls':0,'source_documents_modified':False,'cache_location':rel(CACHE)+' — cache local exclu de publication'},'warnings':warnings,'documents_non_ou_partiellement_indexes':nonindexed,'documents_truncated':truncated,'excluded_files':excluded,'duplicate_groups':duplicates,'redactions_without_values':redactions,'visual_sample_checks':checks,'email_json_extraction':mail_audit,'ocr_preservation_check':{'regression_detected':False,'previous_ocr_documents_preserved':len(preserved_ocr),'method':'Comparaison des pages OCR par SHA avant écriture; arrêt si une page OCR antérieure perd son texte.'}}
    atomic_text(COVER,json.dumps(coverage,ensure_ascii=False,indent=2))
    md=['# Couverture de l’extraction du dossier maître','',f'Générée le {now}. Aucun original ni paquet de remise modifié.','',f"{stats['source_files_seen']} fichiers recensés; {stats['documents_unique']} contenus distincts sélectionnés, {stats['duplicate_paths']} chemins doublons, {stats['excluded_files']} exclusions consignées.",f"{stats['documents_with_text']} documents disposent de texte indexé; {stats['documents_without_text']} restent sans texte. {stats['pages_without_text']} pages/unités sans texte exploitable. {stats['documents_truncated']} documents portent une limite explicite. Corpus : {stats['corpus_bytes']/1_000_000:.2f} Mo.",'','## Méthodes et limites','']+['- '+w for w in warnings]+['',f'OCR local : {ocr_state}. Aucun service externe utilisé.','',f'Plafond de texte par document : {args.max_text_chars:,} caractères après masquage. Toute coupure est marquée sur la page/unité et dans le statut.','', '## Documents sans texte ou avec pages non indexées','', '| Document | Statut | Pages/unités sans texte |','|---|---|---|']
    for d in nonindexed:md.append('| '+d['path'].replace('|','/')+' | '+d['status']+' | '+', '.join(map(str,d['pages_without_text']))+' |')
    md+=['','## Courriels extraits des exports JSON','',f"{mail_audit['unique_messages']} messages uniques reconnus : {mail_audit['messages_added']} corps ajoutés, {mail_audit['messages_already_complete']} corps complets déjà indexés et reliés à leur export source. Les messages ajoutés sont regroupés dans {mail_audit['source_documents_added']} entrées de fichiers sources, une unité par courriel.",'',mail_audit['matching_rule'],'','Seuls les en-têtes lisibles, le corps et les noms des pièces jointes sont extraits. Aucun binaire ni base64 n’est indexé; les JSON ChatGPT, recherches, inventaires techniques et métadonnées de transport restent exclus. La couverture JSON fournit la correspondance de tous les identifiants et de leurs SHA sources.']
    md+=['','## Exclusions et doublons','', 'La liste complète des chemins exclus et de leur motif, ainsi que tous les groupes SHA et alias, figure dans COUVERTURE-EXTRACTION.json. Les annexes identiques renvoient à leur original canonique; aucun document original n’est supprimé.','', '## Contrôle visuel','', 'Aucune validation visuelle exhaustive n’est revendiquée par le script. Le contrôle ciblé effectué séparément doit être inscrit dans visual_sample_checks; toutes les pages OCR restent qualifiées « à contrôler ».']
    for check in checks:
        md+=['',f"### {check.get('title',check['canonical_path'])}",'',f"Pièce : {check['canonical_path']} — page/unité {check['page']}. SHA-256 : {check['sha256']}.",'',check.get('findings',''),'',check.get('method','')]
    atomic_text(COVER.with_suffix('.md'),'\n'.join(md)+'\n')
    print(json.dumps(stats,ensure_ascii=False),flush=True)

if __name__=='__main__':main()
